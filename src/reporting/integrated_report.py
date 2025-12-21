"""
Integrated Excel Report Writer

This module provides a unified report generation system that combines:
- Regressor predictions and accuracy metrics
- Backtest actual trading results
- Benchmark comparisons

Output: Single Excel file with 5 sheets
- Sheet 1: Summary (overall performance)
- Sheet 2: Regressor Metrics (prediction accuracy)
- Sheet 3: Backtest Performance (actual returns)
- Sheet 4: Detailed Trades (stock-by-stock details)
- Sheet 5: Benchmark Comparison (vs SPY, QQQ, etc.)

Author: Quant Trading Team
Date: 2025-12-21
"""

import pandas as pd
import numpy as np
from pathlib import Path
from datetime import datetime
from typing import Dict, List, Optional, Tuple
import logging


class IntegratedReportWriter:
    """
    Integrated Excel report writer for ML backtest results.

    This class manages the creation of a comprehensive Excel report
    that combines regressor predictions with backtest performance.

    Usage:
    ------
    # Create writer
    writer = IntegratedReportWriter()

    # Add data from regressor
    writer.add_sheet('Regressor Metrics', regressor_metrics_df)

    # Add data from ml_backtest
    writer.add_sheet('Backtest Performance', backtest_results_df)

    # Write to file
    filepath = writer.write()
    """

    def __init__(self, output_dir: str = "outputs/reports"):
        """
        Initialize report writer.

        Parameters:
        -----------
        output_dir : str
            Directory to save the report (default: outputs/reports)
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        # Timestamp for filename
        self.timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.filename = f"integrated_report_{self.timestamp}.xlsx"
        self.filepath = self.output_dir / self.filename

        # Data storage for each sheet
        self.sheets = {
            'Summary': None,
            'Regressor Metrics': None,
            'Backtest Performance': None,
            'Detailed Trades': None,
            'Benchmark Comparison': None
        }

        self.logger = logging.getLogger('IntegratedReport')

    def add_sheet(self, sheet_name: str, data: pd.DataFrame):
        """
        Add data for a specific sheet.

        Parameters:
        -----------
        sheet_name : str
            Name of the sheet (must be one of the predefined sheets)
        data : pd.DataFrame
            Data to add to the sheet

        Raises:
        -------
        ValueError
            If sheet_name is not recognized
        """
        if sheet_name not in self.sheets:
            raise ValueError(
                f"Invalid sheet name: {sheet_name}. "
                f"Valid names: {list(self.sheets.keys())}"
            )

        if data is not None and not data.empty:
            self.sheets[sheet_name] = data
            self.logger.info(f"✅ Added sheet: {sheet_name} ({len(data)} rows)")
        else:
            self.logger.warning(f"⚠️  Sheet {sheet_name} is empty, skipping")

    def write(self) -> Path:
        """
        Write all sheets to Excel file.

        Returns:
        --------
        Path
            Path to the created Excel file
        """
        # Check if we have at least one non-empty sheet
        non_empty_sheets = {k: v for k, v in self.sheets.items() if v is not None}

        if not non_empty_sheets:
            self.logger.error("❌ No data to write! All sheets are empty.")
            raise ValueError("Cannot write report: no data provided")

        self.logger.info(f"\n📊 Writing integrated report with {len(non_empty_sheets)} sheets...")

        # Write to Excel
        with pd.ExcelWriter(self.filepath, engine='openpyxl') as writer:
            for sheet_name, df in self.sheets.items():
                if df is not None and not df.empty:
                    # Write data
                    df.to_excel(writer, sheet_name=sheet_name, index=False)

                    # Apply formatting
                    worksheet = writer.sheets[sheet_name]
                    self._format_sheet(worksheet, df, sheet_name)

                    self.logger.info(f"   ✅ {sheet_name}: {len(df)} rows × {len(df.columns)} cols")

        self.logger.info(f"\n✅ Integrated report saved: {self.filepath}")
        return self.filepath

    def _format_sheet(self, worksheet, df: pd.DataFrame, sheet_name: str):
        """
        Apply formatting to Excel sheet.

        Parameters:
        -----------
        worksheet : openpyxl.worksheet.worksheet.Worksheet
            Excel worksheet object
        df : pd.DataFrame
            Dataframe that was written to the sheet
        sheet_name : str
            Name of the sheet (for context-specific formatting)
        """
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        # Header formatting
        header_fill = PatternFill(
            start_color="366092",
            end_color="366092",
            fill_type="solid"
        )
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal='center', vertical='center')

        # Apply header styles
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        # Column width auto-adjustment
        for column in worksheet.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            # Set width (max 50 chars)
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width

        # Freeze header row
        worksheet.freeze_panes = 'A2'

        # Add borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        for row in worksheet.iter_rows():
            for cell in row:
                cell.border = thin_border


def create_summary_sheet(
    regressor_metrics: pd.DataFrame,
    backtest_performance: pd.DataFrame
) -> pd.DataFrame:
    """
    Create Sheet 1: Summary.

    This sheet provides high-level overview of:
    - Total return per period
    - Overall performance metrics
    - Comparison across different periods

    Parameters:
    -----------
    regressor_metrics : pd.DataFrame
        Metrics from regressor evaluation
    backtest_performance : pd.DataFrame
        Performance results from ml_backtest

    Returns:
    --------
    pd.DataFrame
        Summary statistics
    """
    summary = []

    if backtest_performance.empty:
        return pd.DataFrame()

    # Group by period if 'period' column exists
    if 'period' in backtest_performance.columns:
        periods = backtest_performance['period'].unique()
    else:
        periods = ['Overall']
        backtest_performance['period'] = 'Overall'

    # Calculate metrics for each period
    for period in periods:
        period_data = backtest_performance[backtest_performance['period'] == period]

        if 'avg_return' in period_data.columns:
            returns = period_data['avg_return']
        elif 'period_return' in period_data.columns:
            returns = period_data['period_return']
        else:
            continue

        total_return = (1 + returns).prod() - 1
        avg_return = returns.mean()
        std_return = returns.std()

        # Sharpe ratio (assuming quarterly returns)
        sharpe = (avg_return / std_return) * np.sqrt(4) if std_return > 0 else 0.0

        # Maximum drawdown
        cumulative = (1 + returns).cumprod()
        running_max = cumulative.expanding().max()
        drawdown = (cumulative - running_max) / running_max
        max_drawdown = drawdown.min()

        # Win rate
        win_rate = (returns > 0).sum() / len(returns) if len(returns) > 0 else 0.0

        summary.append({
            'Period': period,
            'Total Return': f"{total_return*100:.2f}%",
            'Avg Period Return': f"{avg_return*100:.2f}%",
            'Std Dev': f"{std_return*100:.2f}%",
            'Sharpe Ratio': f"{sharpe:.2f}",
            'Max Drawdown': f"{max_drawdown*100:.2f}%",
            'Win Rate': f"{win_rate*100:.1f}%",
            'Num Periods': len(returns)
        })

    return pd.DataFrame(summary)


def create_regressor_metrics_sheet(predictions_history: List[Dict]) -> pd.DataFrame:
    """
    Create Sheet 2: Regressor Metrics.

    This sheet shows prediction accuracy metrics for each rebalancing period:
    - RMSE, MAE, R² (regression metrics)
    - Accuracy, Precision, Recall (classification metrics)

    Parameters:
    -----------
    predictions_history : List[Dict]
        List of prediction results per period
        Each dict should have: rebalance_date, actual_returns, predicted_returns, etc.

    Returns:
    --------
    pd.DataFrame
        Regressor performance metrics per period
    """
    if not predictions_history:
        return pd.DataFrame()

    from sklearn.metrics import (
        mean_squared_error,
        mean_absolute_error,
        r2_score,
        accuracy_score,
        precision_score,
        recall_score
    )

    metrics = []

    for pred_info in predictions_history:
        try:
            date = pred_info['rebalance_date']
            y_true = pred_info['actual_returns']
            y_pred = pred_info['predicted_returns']

            # Check if we have binary labels
            has_binary = ('actual_labels' in pred_info and
                         'predicted_labels' in pred_info)

            if has_binary:
                y_true_binary = pred_info['actual_labels']
                y_pred_binary = pred_info['predicted_labels']

            # Regression metrics
            rmse = mean_squared_error(y_true, y_pred, squared=False)
            mae = mean_absolute_error(y_true, y_pred)
            r2 = r2_score(y_true, y_pred)

            # Classification metrics (if available)
            if has_binary:
                accuracy = accuracy_score(y_true_binary, y_pred_binary)
                precision = precision_score(y_true_binary, y_pred_binary, zero_division=0)
                recall = recall_score(y_true_binary, y_pred_binary, zero_division=0)
            else:
                # Fallback: use sign of returns
                accuracy = accuracy_score((y_true > 0).astype(int), (y_pred > 0).astype(int))
                precision = precision_score((y_true > 0).astype(int), (y_pred > 0).astype(int), zero_division=0)
                recall = recall_score((y_true > 0).astype(int), (y_pred > 0).astype(int), zero_division=0)

            metrics.append({
                'Period': date.strftime('%Y-%m-%d') if hasattr(date, 'strftime') else str(date),
                'RMSE': f"{rmse:.4f}",
                'MAE': f"{mae:.4f}",
                'R²': f"{r2:.4f}",
                'Accuracy': f"{accuracy*100:.2f}%",
                'Precision': f"{precision*100:.2f}%",
                'Recall': f"{recall*100:.2f}%",
                'Num Predictions': len(y_true)
            })
        except Exception as e:
            logging.warning(f"⚠️  Failed to calculate metrics for {pred_info.get('rebalance_date')}: {e}")
            continue

    return pd.DataFrame(metrics)


def create_backtest_performance_sheet(backtest_results: pd.DataFrame) -> pd.DataFrame:
    """
    Create Sheet 3: Backtest Performance.

    This sheet shows actual trading results per rebalancing period:
    - Period returns
    - Cumulative returns
    - Trading dates

    Parameters:
    -----------
    backtest_results : pd.DataFrame
        Backtest results from ml_backtest.run()

    Returns:
    --------
    pd.DataFrame
        Formatted backtest performance data
    """
    if backtest_results.empty:
        return pd.DataFrame()

    # Format for display
    df = backtest_results.copy()

    # Format dates
    for col in ['rebalance_date', 'actual_buy_date', 'actual_sell_date']:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col]).dt.strftime('%Y-%m-%d')

    # Format percentages
    if 'avg_return' in df.columns:
        df['Period Return'] = (df['avg_return'] * 100).apply(lambda x: f"{x:.2f}%")

        # Calculate cumulative
        df['Cumulative Return'] = ((1 + df['avg_return']).cumprod() - 1) * 100
        df['Cumulative Return'] = df['Cumulative Return'].apply(lambda x: f"{x:.2f}%")

    # Select and rename columns
    display_cols = {
        'rebalance_date': 'Rebalance Date',
        'actual_buy_date': 'Buy Date',
        'actual_sell_date': 'Sell Date',
        'num_stocks': 'Stocks Selected',
        'Period Return': 'Period Return',
        'Cumulative Return': 'Cumulative Return',
        'retrained': 'Model Retrained'
    }

    available_cols = [col for col in display_cols.keys() if col in df.columns]
    result = df[available_cols].rename(columns=display_cols)

    return result


# Export public functions
__all__ = [
    'IntegratedReportWriter',
    'create_summary_sheet',
    'create_regressor_metrics_sheet',
    'create_backtest_performance_sheet'
]
